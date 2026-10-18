"""The /admin/gst page: access, quarter choice, saving, the filed-and-frozen tracker, tax settings, Excel export."""
import io
import json
from datetime import date, datetime

import pytest
from openpyxl import load_workbook

import app as appmod
import billing_client
import sales_stats
from conftest import ADMIN
from test_gst_report import SHOP_RAW

TODAY = date(2026, 10, 18)           # in the quarter after Q2 (Jul-Sep), which is the one being filed


@pytest.fixture(autouse=True)
def _today(monkeypatch):
    monkeypatch.setattr(sales_stats, 'ist_today', lambda now_utc=None: TODAY)


@pytest.fixture()
def site(client, monkeypatch):
    """A product, the real August order (paid and shipped) with its invoice, and an admin session."""
    p = appmod.Product(name='Garam Masala - 50g', description='x', price='55', stock=50)
    appmod.db.session.add(p)
    appmod.db.session.flush()
    o = appmod.SpiceOrder(order_number='HS-24284FEE', user_email='b@example.com', full_name='Buyer', phone='9' * 10, address='a',
                          city='Sindewahi', state='Maharashtra', pincode='441222', subtotal=11000, shipping_cost=3700,
                          total_amount=14700, razorpay_order_id='order_X', razorpay_payment_id='pay_X', payment_status='paid',
                          shipping_status='shipped', created_at=datetime(2026, 8, 29, 9, 53))
    appmod.db.session.add(o)
    appmod.db.session.flush()
    appmod.db.session.add(appmod.OrderItem(order_id=o.id, product_id=p.id, product_name=p.name, quantity=2, unit_price=5500))
    appmod.db.session.commit()
    monkeypatch.setattr(billing_client, 'configured', lambda: False)
    with client.session_transaction() as sess:
        sess['user'] = {'email': ADMIN, 'name': 'Owner'}
    return client


def _filing(fy=2026, q=2):
    row = appmod.GstFiling.query.filter_by(fy_year=fy, quarter=q).first()
    return json.loads(row.data) if row else None


def _post(c, path, **data):
    data.setdefault('fy', '2026')
    data.setdefault('q', '2')
    return c.post(path, data=data, follow_redirects=False)


# ---------- access ----------

@pytest.mark.parametrize('method,path', [('get', '/admin/gst'), ('get', '/admin/gst/export.xlsx'), ('post', '/admin/gst/inputs'),
                                         ('post', '/admin/gst/status'), ('post', '/admin/gst/settings')])
def test_only_admins_can_use_any_gst_route(client, method, path):
    assert getattr(client, method)(path).status_code == 403
    with client.session_transaction() as sess:
        sess['user'] = {'email': 'stranger@example.com', 'name': 'X'}
    assert getattr(client, method)(path).status_code == 403


# ---------- the page ----------

def test_the_page_opens_on_the_quarter_being_filed_and_shows_the_real_order(site):
    html = site.get('/admin/gst').get_data(as_text=True)
    assert 'Q2 FY 2026-27 (Jul-Sep 2026)' in html
    assert 'HS-24284FEE' in html and '₹140.00' in html and '₹7.00' in html          # taxable and GST of the August order
    assert 'NAHID PARWIN AKIL SHEIKH' in html and '27KYIPS1412N1ZT' in html
    assert 'not connected' in html                                                    # the shop is not linked in this test
    assert 'no GST invoice' in html and 'HSN 0910 / 5%' in html                       # both honest warnings
    assert '27-Maharashtra' in html


def test_due_dates_are_flagged_overdue_or_soon(site, monkeypatch):
    html = site.get('/admin/gst?fy=2026&q=2').get_data(as_text=True)
    assert 'GSTR-1' in html and '5 days overdue' in html and '4 days left' in html        # 13 Oct passed, 22 Oct ahead
    monkeypatch.setattr(sales_stats, 'ist_today', lambda now_utc=None: date(2026, 10, 22))
    assert 'due today' in site.get('/admin/gst?fy=2026&q=2').get_data(as_text=True)


def test_bad_or_missing_quarter_falls_back_to_the_latest_finished_one(site):
    for q in ('', '?fy=abc&q=9', '?fy=1900&q=2', '?fy=2026&q=0'):
        assert 'Q2 FY 2026-27' in site.get('/admin/gst' + q).get_data(as_text=True)


def test_other_quarters_can_be_opened(site):
    html = site.get('/admin/gst?fy=2026&q=1').get_data(as_text=True)
    assert 'Q1 FY 2026-27 (Apr-Jun 2026)' in html and 'No paid website orders in this quarter' in html


def test_shop_figures_are_combined_when_connected(site, monkeypatch):
    monkeypatch.setattr(billing_client, 'configured', lambda: True)
    monkeypatch.setattr(billing_client, 'get_gst_summary', lambda s, e: SHOP_RAW)
    html = site.get('/admin/gst?fy=2026&q=2').get_data(as_text=True)
    assert 'HS/26-27/A-0001' in html and '27ABCDE1234F1Z5' in html                    # the shop's B2B invoice
    assert '₹1,040.00' in html                                                       # 14,000 online + 90,000 shop (paise) = 1,040.00 taxable
    assert 'Spice Co' in html and 'not connected' not in html


def test_shop_problems_become_a_clear_do_not_file_warning(site, monkeypatch):
    monkeypatch.setattr(billing_client, 'configured', lambda: True)
    def boom(s, e):
        raise billing_client.BillingUnavailable('The shop database could not be reached. Try again in a few minutes.')
    monkeypatch.setattr(billing_client, 'get_gst_summary', boom)
    html = site.get('/admin/gst?fy=2026&q=2').get_data(as_text=True)
    assert 'could not be loaded' in html and 'do not file from these totals' in html and 'HS-24284FEE' in html


# ---------- saving the editable figures ----------

def test_saving_input_credit_late_fee_and_notes(site):
    r = _post(site, '/admin/gst/inputs', itc_igst='0', itc_cgst='1,000.50', itc_sgst='1000.50', late_fee_cgst='250', late_fee_sgst='250',
              interest_igst='', interest_cgst='', interest_sgst='', notes='CA to check supplier X')
    assert r.status_code == 302
    f = _filing()
    assert f['inputs']['itc_cgst'] == 100050 and f['inputs']['late_fee_sgst'] == 25000 and f['inputs']['interest_cgst'] == ''
    html = site.get('/admin/gst?fy=2026&q=2').get_data(as_text=True)
    assert 'value="1000.50"' in html and 'CA to check supplier X' in html
    assert '₹500.00' in html                                                           # late fee 250 + 250 now in total cash


@pytest.mark.parametrize('bad', ['abc', '-5', '1e999', '12,3a'])
def test_invalid_amounts_are_refused_and_nothing_is_saved(site, bad):
    _post(site, '/admin/gst/inputs', itc_cgst=bad)
    assert _filing() is None


# ---------- the tracker: draft, review, filed and frozen ----------

def test_status_moves_through_draft_and_review(site):
    _post(site, '/admin/gst/status', action='review')
    assert _filing()['status'] == 'review'
    _post(site, '/admin/gst/status', action='draft')
    assert _filing()['status'] == 'draft'


def test_unknown_actions_are_rejected(site):
    assert _post(site, '/admin/gst/status', action='destroy').status_code == 400


def test_marking_as_filed_freezes_the_figures_and_records_the_references(site):
    _post(site, '/admin/gst/inputs', itc_igst='0', itc_cgst='1.00', itc_sgst='1.00')
    _post(site, '/admin/gst/status', action='filed', gstr1_arn='AD2710266689458', gstr1_date='2026-10-20',
          gstr3b_arn='AD271026999999X', gstr3b_date='2026-10-21', tax_paid='5.00')
    f = _filing()
    assert f['status'] == 'filed' and f['gstr1_arn'] == 'AD2710266689458' and f['tax_paid'] == 500
    snap = f['snapshot']
    assert snap['outward']['taxable'] == 14000 and snap['outward']['tax'] == 700 and snap['itc_claim']['cgst'] == 100
    html = site.get('/admin/gst?fy=2026&q=2').get_data(as_text=True)
    assert 'Filed and frozen' in html and 'AD2710266689458' in html and 'Reopen for an amendment' in html


def test_a_filed_quarter_cannot_be_edited_or_re_filed(site):
    _post(site, '/admin/gst/status', action='filed')
    before = _filing()
    _post(site, '/admin/gst/inputs', itc_cgst='999')
    _post(site, '/admin/gst/status', action='draft')
    _post(site, '/admin/gst/status', action='filed', gstr1_arn='CHANGED')
    assert _filing() == before and _filing()['status'] == 'filed'


def test_frozen_figures_survive_later_changes_and_the_page_warns_about_drift(site):
    _post(site, '/admin/gst/status', action='filed')
    frozen = _filing()['snapshot']['outward']['taxable']
    o = appmod.SpiceOrder.query.one()
    o.shipping_status = 'cancelled'                                    # the sale is later cancelled
    appmod.db.session.commit()
    assert _filing()['snapshot']['outward']['taxable'] == frozen == 14000
    assert 'changed since filing' in site.get('/admin/gst?fy=2026&q=2').get_data(as_text=True)


def test_reopening_keeps_the_old_filing_in_history(site):
    _post(site, '/admin/gst/status', action='filed', gstr3b_arn='ARN1')
    _post(site, '/admin/gst/status', action='reopen')
    f = _filing()
    assert f['status'] == 'review' and f['snapshot'] is None and len(f['history']) == 1
    assert f['history'][0]['gstr3b_arn'] == 'ARN1' and f['history'][0]['snapshot']['outward']['taxable'] == 14000
    _post(site, '/admin/gst/inputs', itc_cgst='5')                      # editable again
    assert _filing()['inputs']['itc_cgst'] == 500


def test_bad_dates_and_amounts_in_the_tracker_are_refused(site):
    _post(site, '/admin/gst/status', action='filed', gstr1_date='22/10/2026')
    assert _filing() is None
    _post(site, '/admin/gst/status', action='filed', tax_paid='lots')
    assert _filing() is None


def test_each_quarter_has_its_own_record(site):
    _post(site, '/admin/gst/status', action='review', q='2')
    _post(site, '/admin/gst/status', action='filed', q='3')
    assert _filing(q=2)['status'] == 'review' and _filing(q=3)['status'] == 'filed'
    assert appmod.GstFiling.query.count() == 2


# ---------- tax settings ----------

def test_saving_hsn_rate_and_confirmation_per_product(site):
    pid = appmod.Product.query.one().id
    r = _post(site, '/admin/gst/settings', **{f'hsn_{pid}': '0910 99', f'rate_{pid}': '12', f'confirmed_{pid}': 'on',
                                              'shipping_rate': '18', 'gstin': '27kyips1412n1zt', 'address': 'New address'})
    assert r.status_code == 302
    cfg = appmod.gst_product_tax()
    assert cfg[str(pid)] == {'hsn': '091099', 'rate': 12, 'confirmed': True}
    assert appmod.gst_shipping_rate() == 18 and appmod.gst_seller()['gstin'] == '27KYIPS1412N1ZT'
    assert appmod.gst_seller()['address'] == 'New address'
    assert 'HSN 0910 / 5%' not in site.get('/admin/gst?fy=2026&q=2').get_data(as_text=True)      # no longer an unconfirmed default


@pytest.mark.parametrize('field,value', [('rate', 'abc'), ('rate', '99'), ('rate', '-1'), ('hsn', ''), ('ship', 'x'), ('gstin', 'BADGSTIN')])
def test_invalid_tax_settings_are_refused_and_nothing_changes(site, field, value):
    pid = appmod.Product.query.one().id
    data = {f'hsn_{pid}': '0910', f'rate_{pid}': '5', 'shipping_rate': '5', 'gstin': ''}
    data.update({'rate': {f'rate_{pid}': value}, 'hsn': {f'hsn_{pid}': value}, 'ship': {'shipping_rate': value},
                 'gstin': {'gstin': value}}[field])
    _post(site, '/admin/gst/settings', **data)
    assert appmod.gst_get('product_tax', '') == '' and appmod.gst_get('gstin', '') == ''


def test_a_rate_change_is_used_for_new_invoices_only(site):
    pid = appmod.Product.query.one().id
    _post(site, '/admin/gst/settings', **{f'hsn_{pid}': '2103', f'rate_{pid}': '12', f'confirmed_{pid}': 'on', 'shipping_rate': '12'})
    order = appmod.SpiceOrder.query.one()
    appmod.db.session.add(appmod.SpiceOrder(order_number='HS-NEW', user_email='n@example.com', full_name='N', phone='9' * 10, address='a',
                                            city='Pune', state='Maharashtra', pincode='411001', subtotal=11000, shipping_cost=3700,
                                            total_amount=14700, razorpay_order_id='order_N', payment_status='pending'))
    appmod.db.session.commit()
    new = appmod.SpiceOrder.query.filter_by(order_number='HS-NEW').one()
    appmod.db.session.add(appmod.OrderItem(order_id=new.id, product_id=pid, product_name='Garam Masala - 50g', quantity=2, unit_price=5500))
    appmod.db.session.commit()
    appmod.finalize_paid_order('order_N', 'pay_N')
    lines = appmod.GstInvoice.query.filter_by(order_id=new.id).one().doc['breakup']['lines']
    assert lines[0]['rate'] == 12 and lines[0]['hsn'] == '2103'
    assert order.id != new.id and appmod.GstInvoice.query.filter_by(order_id=order.id).first() is None


# ---------- the Excel file ----------

def _book(c, q='2'):
    r = c.get(f'/admin/gst/export.xlsx?fy=2026&q={q}')
    assert r.status_code == 200
    assert 'spreadsheetml' in r.mimetype and f'GST-FY2026-27-Q{q}.xlsx' in r.headers['Content-Disposition']
    return load_workbook(io.BytesIO(r.data))


def test_excel_has_every_sheet_and_the_right_numbers(site):
    wb = _book(site)
    assert wb.sheetnames == ['GSTR-3B', 'GSTR-1 B2CS', 'GSTR-1 B2B', 'GSTR-1 HSN', 'Credit notes', 'Documents issued',
                             'Online orders', 'Purchases (ITC)', 'Read me']
    b2cs = wb['GSTR-1 B2CS']
    assert b2cs['A5'].value == '27-Maharashtra' and b2cs['B5'].value == 5 and b2cs['C5'].value == 140.0      # rupees, not paise
    assert b2cs['E5'].value == 3.5 and b2cs['F5'].value == 3.5 and b2cs['G5'].value == 140.0 and b2cs['H5'].value == 0
    orders = wb['Online orders']
    assert orders['A5'].value == 'HS-24284FEE' and orders['F5'].value == 140.0 and orders['J5'].value == 147.0
    hsn = wb['GSTR-1 HSN']
    assert hsn['B5'].value == '0910' and hsn['E5'].value == 2 and hsn['F5'].value == 140.0


def test_excel_summary_has_live_formulas_and_yellow_input_cells(site):
    ws = _book(site)['GSTR-3B']
    assert ws['B7'].value == '=SUM(B5:B6)' and ws['D7'].value == '=SUM(D5:D6)'
    assert ws['B5'].value == 140.0 and ws['D5'].value == 3.5 and ws['E5'].value == 3.5 and ws['B6'].value == 0
    assert ws['C16'].value == '=MAX(0,C14-C15)' and ws['C25'].value == '=SUM(C19:E19)+SUM(C22:E23)'
    for cell in ('C11', 'D11', 'E11', 'D22', 'E22', 'C23', 'D23', 'E23'):                    # the cells the owner may edit
        assert ws[cell].fill.fgColor.rgb.endswith('FFF2B3'), cell
    assert 'Not a return' in ws['A28'].value


def test_excel_uses_the_saved_figures_and_marks_missing_shop_data(site):
    _post(site, '/admin/gst/inputs', itc_cgst='1.25', itc_sgst='1.25', itc_igst='0', late_fee_cgst='250', late_fee_sgst='250')
    wb = _book(site)
    ws = wb['GSTR-3B']
    assert ws['D11'].value == 1.25 and ws['D22'].value == 250 and ws['E22'].value == 250
    assert any('shop figures are missing' in str(c.value) for c in wb['Read me']['A'])
    assert any('no GST invoice' in str(c.value) for c in wb['Read me']['A'])


def test_excel_includes_the_shops_tables_when_connected(site, monkeypatch):
    monkeypatch.setattr(billing_client, 'configured', lambda: True)
    monkeypatch.setattr(billing_client, 'get_gst_summary', lambda s, e: SHOP_RAW)
    wb = _book(site)
    assert wb['GSTR-1 B2B']['A5'].value == 'HS/26-27/A-0001' and wb['GSTR-1 B2B']['C5'].value == '27ABCDE1234F1Z5'
    assert wb['Credit notes']['A5'].value == 'HS/26-27/CN-1' and wb['Purchases (ITC)']['B5'].value == 'Spice Co'
    assert wb['GSTR-3B']['B6'].value == 900.0 and wb['GSTR-3B']['B8'].value == 70.0               # shop net taxable; bill of supply
    assert wb['GSTR-1 B2CS']['C5'].value == 140.0 + 500.0


def test_excel_works_for_an_empty_quarter(site):
    wb = _book(site, q='1')
    assert wb['Online orders']['A5'].value == 'Total' and wb['GSTR-3B']['B7'].value == '=SUM(B5:B6)'
